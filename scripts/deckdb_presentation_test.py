"""Presentation test of the deckdb-mode charts (order of 2026-10-21 g): render every chart at 1920x1080
(figsize 12.8 x 7.2 in at 150 dpi) into a scratch directory, assemble them into presentation_test.pptx (one
16:9 slide per chart, the image filling the slide) and report readability per chart:

  - smallest font in slide points: the image is scaled to the 13.333 in slide width, so a font of f pt at
    150 dpi in a 1920 px wide image is f * (13.333 * 72 / 1920) * (150 / 72) = ~1.04 f pt on the slide;
  - text outside the figure (cut off) and overlapping tick / annotation labels (bounding boxes, same axis);
  - number of rows / bars / lines.

The materials themselves are not changed (deckdb_materials.RENDER_OVERRIDE is used only here).

    python scripts/deckdb_presentation_test.py --as-of 2026-10-21 --out outputs/deckdb_working/presentation_test
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
from matplotlib.figure import Figure  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import deckdb_input as di  # noqa: E402
import deckdb_materials as dm  # noqa: E402

SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
W_PX, DPI = 1920, 150
MIN_SLIDE_PT = 14.0
REPORT: dict = {}


def _analyse(fig: Figure, path: str) -> None:
    renderer = fig.canvas.get_renderer()
    fw, fh = fig.bbox.width, fig.bbox.height
    sizes, cut, overlaps, rows = [], 0, 0, 0
    texts = [t for t in fig.findobj(matplotlib.text.Text) if t.get_visible() and t.get_text().strip()]
    for t in texts:
        sizes.append(t.get_fontsize())
        bb = t.get_window_extent(renderer)
        if bb.x0 < -1 or bb.y0 < -1 or bb.x1 > fw + 1 or bb.y1 > fh + 1:
            cut += 1
    for ax in fig.axes:
        for labels in (ax.get_xticklabels(), ax.get_yticklabels()):
            boxes = [lb.get_window_extent(renderer) for lb in labels if lb.get_visible() and lb.get_text().strip()]
            rows = max(rows, len(boxes))
            overlaps += sum(1 for i in range(len(boxes) - 1) if boxes[i].overlaps(boxes[i + 1]))
        ann = [t.get_window_extent(renderer) for t in ax.texts if t.get_visible() and t.get_text().strip()]
        overlaps += sum(1 for i in range(len(ann)) for j in range(i + 1, len(ann)) if ann[i].overlaps(ann[j]))
    scale = SLIDE_W_IN * 72 / (fig.get_figwidth() * fig.dpi) * (fig.dpi / 72)  # slide pt per figure pt
    REPORT[path] = {
        "min_font_slide_pt": round(min(sizes) * scale, 1) if sizes else None,
        "texts_below_14pt": sum(1 for s in sizes if s * scale < MIN_SLIDE_PT),
        "texts": len(sizes),
        "cut_texts": cut,
        "overlapping_labels": overlaps,
        "rows_or_bars": rows,
    }


_orig_savefig = Figure.savefig


def _savefig(self, fname, *args, **kwargs):
    self.canvas.draw()
    _analyse(self, str(fname))
    return _orig_savefig(self, fname, *args, **kwargs)


def build_pptx(pngs: list[Path], out: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
    for png in pngs:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--as-of", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--trend-weeks", type=int, default=12)
    args = ap.parse_args()
    out = Path(args.out)
    cfg = di.load_deckdb_config()
    export = di.load_export(cfg.export_path)
    dm.RENDER_OVERRIDE.update({"figsize": (W_PX / DPI, 1080 / DPI), "dpi": DPI})
    Figure.savefig = _savefig
    try:
        files = dm.build_deckdb_materials(export, cfg, date.fromisoformat(args.as_of), out, args.trend_weeks,
                                          log=lambda _m: None)
    finally:
        Figure.savefig = _orig_savefig
        dm.RENDER_OVERRIDE.clear()
    pngs = [f for f in files if f.suffix == ".png"]
    pptx = out / "presentation_test.pptx"
    build_pptx(pngs, pptx)
    from PIL import Image

    print(f"{len(pngs)} charts -> {pptx}")
    print("| chart | px | min font (slide pt) | texts < 14 pt | cut texts | overlapping labels | rows/bars |")
    print("|---|---|---|---|---|---|---|")
    for p in pngs:
        r = next((v for k, v in REPORT.items() if Path(k).name == p.name and Path(k).parent.name == p.parent.name), {})
        w, h = Image.open(p).size
        print(f"| {p.parent.name}/{p.name} | {w}x{h} | {r.get('min_font_slide_pt')} | {r.get('texts_below_14pt')}/"
              f"{r.get('texts')} | {r.get('cut_texts')} | {r.get('overlapping_labels')} | {r.get('rows_or_bars')} |")
    return 0


if __name__ == "__main__":
    sys.exit(main())
